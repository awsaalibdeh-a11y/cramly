"""Cramly's server without the internet: the AI is faked."""

import datetime as dt
import io
import json
import os
import time
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["DATABASE_URL"] = f"sqlite:///{os.path.join(tempfile.mkdtemp(), 'test.db')}"
os.environ["OPENAI_API_KEY"] = "test"

import ai  # noqa: E402
import app as server  # noqa: E402
import ingest  # noqa: E402

PLAN = {"title": "Cells", "emoji": "🔬", "topics": [
    {"title": "Cell basics", "summary": "What cells are", "points": ["nucleus", "membrane"]},
    {"title": "Organelles", "summary": "Parts of a cell", "points": ["mitochondria", "ribosome"]},
    {"title": "Cell division", "summary": "Mitosis", "points": ["mitosis"]}]}
QUIZ = [{"q": f"Question {i}?", "options": ["a", "b", "c", "d"], "answer": i % 4, "why": "because"} for i in range(6)]
CARDS = [{"front": f"term {i}", "back": f"meaning {i}"} for i in range(10)]
NOTES = "Hook line.\n\n## What it is\n- **Cell**: unit of life\n\n## Remember\n- one\n- two"
TEXT = "The cell is the basic unit of life. " * 12


def fake_plan(text, hint="", on_status=None):
    if on_status:
        on_status("Planning your topics…")
    return json.loads(json.dumps(PLAN))


class Base(unittest.TestCase):
    def setUp(self):
        self.c = server.app.test_client()
        server._hits.clear()
        self.key = self.c.post("/api/account", json={"name": "Aws"}).get_json()["key"]
        self.h = {"Authorization": f"Bearer {self.key}", "X-Tz-Offset": "-180"}

    def events(self, resp):
        return [json.loads(x) for x in resp.get_data(as_text=True).splitlines() if x.strip()]

    def make_set(self, text=TEXT):
        with mock.patch.object(ai, "make_plan", fake_plan):
            ev = self.events(self.c.post("/api/sets", headers=self.h, data={"text": text}))
        return next(e["id"] for e in ev if e["type"] == "set")

    def topics(self, sid):
        return self.c.get(f"/api/sets/{sid}", headers=self.h).get_json()["topics"]



class Cramly(Base):
    def test_needs_a_key(self):
        self.assertEqual(self.c.get("/api/sets").status_code, 401)
        self.assertEqual(self.c.get("/api/sets", headers={"Authorization": "Bearer nope"}).status_code, 401)

    def test_create_set_streams_progress_then_builds_topics(self):
        with mock.patch.object(ai, "make_plan", fake_plan):
            ev = self.events(self.c.post("/api/sets", headers=self.h, data={"text": TEXT, "title": ""}))
        self.assertEqual([e["type"] for e in ev if e["type"] != "ping"], ["status", "status", "set", "done"])
        sid = next(e["id"] for e in ev if e["type"] == "set")
        d = self.c.get(f"/api/sets/{sid}", headers=self.h).get_json()
        self.assertEqual((d["set"]["title"], d["set"]["topics"], len(d["materials"])), ("Cells", 3, 1))
        self.assertEqual([t["title"] for t in d["topics"]], ["Cell basics", "Organelles", "Cell division"])

    def test_empty_upload_is_refused(self):
        self.assertEqual(self.c.post("/api/sets", headers=self.h, data={"text": "too short"}).status_code, 400)

    def test_a_friendly_error_is_streamed_when_the_ai_fails(self):
        with mock.patch.object(ai, "make_plan", side_effect=ai.AIError("The AI is busy right now. Try again in a moment.")):
            ev = self.events(self.c.post("/api/sets", headers=self.h, data={"text": TEXT}))
        self.assertEqual(ev[-2], {"type": "error", "text": "The AI is busy right now. Try again in a moment."})
        self.assertEqual(self.c.get("/api/sets", headers=self.h).get_json()["sets"], [])

    def test_sets_are_private_to_their_owner(self):
        sid = self.make_set()
        other = self.c.post("/api/account", json={}).get_json()["key"]
        self.assertEqual(self.c.get(f"/api/sets/{sid}", headers={"Authorization": f"Bearer {other}"}).status_code, 404)
        tid = self.topics(sid)[0]["id"]
        self.assertEqual(self.c.get(f"/api/topics/{tid}/notes", headers={"Authorization": f"Bearer {other}"}).status_code, 404)
        self.assertEqual(self.c.delete(f"/api/sets/{sid}", headers={"Authorization": f"Bearer {other}"}).status_code, 404)

    def test_notes_cards_quiz_are_made_once_then_kept(self):
        tid = self.topics(self.make_set())[0]["id"]
        with mock.patch.object(ai, "make_notes", return_value=NOTES) as notes, mock.patch.object(ai, "make_cards", return_value=CARDS) as cards, \
                mock.patch.object(ai, "make_quiz", return_value=QUIZ) as quiz:
            for _ in range(2):
                self.assertIn("Remember", self.c.get(f"/api/topics/{tid}/notes", headers=self.h).get_json()["notes"])
                self.assertEqual(len(self.c.get(f"/api/topics/{tid}/cards", headers=self.h).get_json()["cards"]), 10)
                self.assertEqual(len(self.c.get(f"/api/topics/{tid}/quiz", headers=self.h).get_json()["questions"]), 6)
        self.assertEqual((notes.call_count, cards.call_count, quiz.call_count), (1, 1, 1))

    def test_mastery_comes_from_quiz_80_percent(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        post = lambda **b: self.c.post(f"/api/topics/{tid}/event", headers=self.h, json=b).get_json()
        self.assertEqual(post(kind="read")["status"], 1)
        self.assertEqual(post(kind="quiz", score=3, total=6)["status"], 1)
        r = post(kind="quiz", score=5, total=6)
        self.assertEqual((r["status"], r["quiz_best"]), (2, 83))
        self.assertEqual(post(kind="quiz", score=1, total=6)["status"], 2)       # never goes back down
        self.assertEqual(self.c.get(f"/api/sets/{sid}", headers=self.h).get_json()["set"]["mastered"], 1)

    def test_flashcard_review_schedules_and_masters(self):
        tid = self.topics(self.make_set())[0]["id"]
        with mock.patch.object(ai, "make_cards", return_value=CARDS[:2]):
            cards = self.c.get(f"/api/topics/{tid}/cards", headers=self.h).get_json()["cards"]
        rv = lambda cid, g: self.c.post(f"/api/cards/{cid}/review", headers=self.h, json={"grade": g}).get_json()
        self.assertEqual(rv(cards[0]["id"], 0)["box"], 0)
        self.assertEqual(rv(cards[0]["id"], 2)["box"], 1)
        self.assertEqual(rv(cards[0]["id"], 3)["box"], 3)
        self.assertEqual(rv(cards[1]["id"], 3)["box"], 2)
        self.assertEqual(rv(cards[1]["id"], 3)["status"], 2)                      # every card in box >= 3: mastered
        due = self.c.get("/api/due", headers=self.h).get_json()["cards"]
        self.assertEqual(due, [])                                                  # all pushed into the future

    def test_chat_streams_and_uses_the_material(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        seen = {}

        def fake_chat(mode, context, messages, style=""):
            seen.update(mode=mode, context=context, style=style)
            yield "Hello "
            yield "there"
        with mock.patch.object(ai, "chat", fake_chat):
            r = self.c.post("/api/chat", headers=self.h, json={"set_id": sid, "topic_id": tid, "mode": "guided", "messages": [{"role": "user", "content": "teach me"}]})
        text = "".join(e.get("text", "") for e in self.events(r) if e["type"] == "delta")
        self.assertEqual(text, "Hello there")
        self.assertEqual((seen["mode"], seen["style"]), ("guided", ""))
        self.assertIn("CURRENT TOPIC: Cell basics", seen["context"])
        self.assertIn("basic unit of life", seen["context"])
        bad = self.c.post("/api/chat", headers=self.h, json={"set_id": sid, "messages": [{"role": "assistant", "content": "x"}]})
        self.assertEqual(bad.status_code, 400)

    def test_exam_date_makes_a_day_by_day_plan(self):
        sid = self.make_set()
        exam = (dt.date.today() + dt.timedelta(days=4)).isoformat()
        self.assertEqual(self.c.patch(f"/api/sets/{sid}", headers=self.h, json={"exam": exam}).status_code, 200)
        d = self.c.get(f"/api/sets/{sid}", headers=self.h).get_json()
        self.assertEqual(len(d["plan"]), 3)
        self.assertTrue(all(p["date"] < exam for p in d["plan"]))                  # the exam day itself stays free
        ev = self.c.get("/api/calendar", headers=self.h).get_json()["events"]
        self.assertEqual([e["kind"] for e in ev].count("exam"), 1)
        self.assertEqual(self.c.patch(f"/api/sets/{sid}", headers=self.h, json={"exam": "banana"}).status_code, 400)

    def test_schedule_spreads_topics_and_skips_mastered(self):
        today = dt.date(2026, 10, 1)
        topics = [{"id": i, "status": 2 if i == 0 else 0} for i in range(7)]
        plan = server.schedule("2026-10-05", topics, today)
        self.assertEqual(len(plan), 6)
        self.assertEqual(max(d for d, _ in plan), dt.date(2026, 10, 3))
        self.assertEqual(server.schedule("2026-09-01", topics, today), [])
        self.assertEqual(server.schedule("", topics, today), [])

    def test_streak_counts_consecutive_days(self):
        sid = self.make_set()
        uid = server.db.one(server.select(server.db.sets.c.user_id).where(server.db.sets.c.id == sid))["user_id"]
        today = (dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=3)).date()           # the test's browser is UTC+3
        for back in (1, 2, 3):
            server.db.run(server.insert(server.db.activity).values(user_id=uid, day=(today - dt.timedelta(days=back)).isoformat(), n=1))
        with server.app.test_request_context(headers={"X-Tz-Offset": "-180"}):
            self.assertGreaterEqual(server.streak({"id": uid}), 4)

    def test_add_material_only_adds_new_topics(self):
        sid = self.make_set()
        more = {**PLAN, "topics": [PLAN["topics"][0], {"title": "Genetics", "summary": "DNA", "points": ["gene"]}]}
        with mock.patch.object(ai, "make_plan", lambda t, h="", on_status=None: json.loads(json.dumps(more))):
            ev = self.events(self.c.post(f"/api/sets/{sid}/materials", headers=self.h, data={"text": TEXT + " Genetics adds DNA and genes."}))
        self.assertIn({"type": "added", "topics": 1}, ev)
        self.assertEqual([t["title"] for t in self.topics(sid)][-1], "Genetics")

    def test_practice_test_mixes_topic_quizzes(self):
        sid = self.make_set()
        with mock.patch.object(ai, "make_quiz", return_value=QUIZ):
            qs = self.c.get(f"/api/sets/{sid}/test", headers=self.h).get_json()["questions"]
        self.assertGreaterEqual(len(qs), 8)
        self.assertEqual({q["topic"] for q in qs} <= {t["title"] for t in PLAN["topics"]}, True)

    def test_pairing_a_second_device(self):
        code = self.c.post("/api/pair", headers=self.h).get_json()["code"]
        key2 = self.c.post("/api/pair/claim", json={"code": code}).get_json()["key"]
        self.assertEqual(self.c.get("/api/me", headers={"Authorization": f"Bearer {key2}"}).get_json()["name"], "Aws")
        self.assertEqual(self.c.post("/api/pair/claim", json={"code": code}).status_code, 400)       # one use

    def test_delete_account_removes_everything(self):
        sid = self.make_set()
        self.assertEqual(self.c.delete("/api/me", headers=self.h).status_code, 200)
        self.assertEqual(self.c.get("/api/me", headers=self.h).status_code, 401)
        for table in (server.db.topics, server.db.materials, server.db.sets):
            col = table.c.set_id if "set_id" in table.c else table.c.id
            self.assertEqual(server.db.many(server.select(table).where(col == sid)), [])    # no orphans left behind

    def test_install_files(self):
        self.assertEqual(self.c.get("/manifest.webmanifest").get_json()["short_name"], "Cramly")
        sw = self.c.get("/sw.js")
        self.assertEqual(sw.headers["Service-Worker-Allowed"], "/")
        sw.close()


SWIPES = [{"q": f"Card {i}?", "options": ["right", "wrong"], "answer": i % 2, "why": "because"} for i in range(14)]


class Round2(Base):
    """Swipe game, explain-it-back and the database upgrade."""

    def test_swipe_cards_are_made_once_and_more_adds_new_ones(self):
        tid = self.topics(self.make_set())[0]["id"]
        batch = iter([SWIPES, [{"q": f"New {i}?", "options": ["a", "b"], "answer": 0, "why": "w"} for i in range(14)] + SWIPES[:2]])
        with mock.patch.object(ai, "make_swipes", side_effect=lambda *a, **k: next(batch)) as made:
            d = self.c.get(f"/api/topics/{tid}/swipes", headers=self.h).get_json()
            self.assertEqual((len(d["cards"]), d["pool"]), (14, 14))
            self.assertEqual(made.call_count, 1)
            self.c.get(f"/api/topics/{tid}/swipes", headers=self.h)
            self.assertEqual(made.call_count, 1)                                    # cached
            d = self.c.get(f"/api/topics/{tid}/swipes?more=1", headers=self.h).get_json()
            self.assertEqual(d["pool"], 28)                                         # 14 new; the 2 repeats were dropped
        for c in d["cards"]:
            self.assertEqual((len(c["options"]), c["answer"] in (0, 1)), (2, True))

    def test_whole_set_swipe_round_mixes_topics(self):
        sid = self.make_set()
        with mock.patch.object(ai, "make_swipes", return_value=SWIPES):
            cards = self.c.get(f"/api/sets/{sid}/swipes", headers=self.h).get_json()["cards"]
        self.assertEqual(len(cards), 14)
        self.assertTrue({c["topic"] for c in cards} <= {t["title"] for t in PLAN["topics"]})
        self.assertGreater(len({c["topic_id"] for c in cards}), 1)

    def test_explain_it_back_grades_and_can_master(self):
        tid = self.topics(self.make_set())[0]["id"]
        verdict = {"score": 91, "verdict": "Great", "got": ["a"], "missing": [], "tip": "t"}
        with mock.patch.object(ai, "grade_explanation", return_value=verdict) as grade:
            short = self.c.post(f"/api/topics/{tid}/explain", headers=self.h, json={"text": "too short"})
            self.assertEqual(short.status_code, 400)
            r = self.c.post(f"/api/topics/{tid}/explain", headers=self.h, json={"text": "Cells are the basic unit of life and they contain a nucleus and organelles."})
        self.assertEqual(r.get_json()["score"], 91)
        self.assertEqual(grade.call_count, 1)
        post = lambda **b: self.c.post(f"/api/topics/{tid}/event", headers=self.h, json=b).get_json()
        self.assertEqual(post(kind="swipe")["status"], 1)
        self.assertEqual(post(kind="explain", score=60)["status"], 1)
        self.assertEqual(post(kind="explain", score=88)["status"], 2)               # explaining it well masters it
        self.assertEqual(self.c.post(f"/api/topics/{tid}/event", headers=self.h, json={"kind": "explain"}).status_code, 400)

    def test_old_databases_get_the_new_column(self):
        from sqlalchemy import create_engine, inspect, text
        old = create_engine("sqlite://")
        with old.begin() as c:
            c.execute(text("CREATE TABLE topics (id INTEGER PRIMARY KEY, title TEXT)"))
        with mock.patch.object(server.db, "engine", old):
            server.db._migrate()
            server.db._migrate()                                                    # safe to run twice
            self.assertIn("swipes", {c["name"] for c in inspect(old).get_columns("topics")})


class StudyTools(Base):
    """The extra study tools."""

    def first_topic(self):
        sid = self.make_set()
        return sid, self.topics(sid)[0]["id"]

    def test_podcast_is_made_once(self):
        sid, tid = self.first_topic()
        ep = {"title": "Cells 101", "lines": [{"host": "AB"[i % 2], "text": f"line {i}"} for i in range(10)]}
        with mock.patch.object(ai, "make_podcast", return_value=ep) as made:
            for _ in range(2):
                d = self.c.get(f"/api/topics/{tid}/podcast", headers=self.h).get_json()
        self.assertEqual((d["title"], len(d["lines"]), d["topic"], made.call_count), ("Cells 101", 10, "Cell basics", 1))
        with mock.patch.object(ai, "make_podcast", return_value={"title": "x", "lines": []}):
            tid2 = self.topics(sid)[1]["id"]
            self.assertEqual(self.c.get(f"/api/topics/{tid2}/podcast", headers=self.h).status_code, 502)    # too short: refused

    def test_cheat_sheet_is_cached_and_can_be_remade(self):
        sid, _ = self.first_topic()
        with mock.patch.object(ai, "make_cheatsheet", side_effect=["## One\n- a", "## Two\n- b"]) as made:
            self.assertIn("One", self.c.get(f"/api/sets/{sid}/cheatsheet", headers=self.h).get_json()["sheet"])
            self.assertIn("One", self.c.get(f"/api/sets/{sid}/cheatsheet", headers=self.h).get_json()["sheet"])
            self.assertEqual(made.call_count, 1)
            self.assertIn("Two", self.c.get(f"/api/sets/{sid}/cheatsheet?fresh=1", headers=self.h).get_json()["sheet"])

    def test_mind_map_comes_straight_from_the_plan(self):
        sid, _ = self.first_topic()
        d = self.c.get(f"/api/sets/{sid}/mindmap", headers=self.h).get_json()
        self.assertEqual((d["title"], [t["title"] for t in d["children"]]), ("Cells", ["Cell basics", "Organelles", "Cell division"]))
        self.assertEqual([p["title"] for p in d["children"][0]["children"]], ["nucleus", "membrane"])

    def test_snap_and_solve_takes_text_or_a_photo(self):
        sid, _ = self.first_topic()
        solved = {"solution": "Step 1 ...", "final": "42"}
        with mock.patch.object(ai, "solve", return_value=solved) as solve:
            self.assertEqual(self.c.post("/api/solve", headers=self.h, data={"text": "hi"}).status_code, 400)       # nothing to solve
            ok = self.c.post("/api/solve", headers=self.h, data={"text": "What is 6 x 7?", "set_id": str(sid)})
            self.assertEqual(ok.get_json()["final"], "42")
            self.assertIn("basic unit of life", solve.call_args[0][3])                                                 # their material came along
            photo = self.c.post("/api/solve", headers=self.h, data={"image": (io.BytesIO(b"\x89PNG fake"), "q.png")})
            self.assertEqual(photo.status_code, 200)
            self.assertEqual(solve.call_args[0][2], "image/png")
            bad = self.c.post("/api/solve", headers=self.h, data={"image": (io.BytesIO(b"MZ"), "virus.exe")})
            self.assertEqual(bad.status_code, 400)

    def test_exam_builder_makes_questions_for_the_chosen_topics(self):
        sid = self.make_set()
        ts = self.topics(sid)
        seen = []

        def fake_quiz(topic, ctx, n=6, level="medium"):
            seen.append((topic["title"], level))
            return [dict(QUIZ[i % 6]) for i in range(n)]
        with mock.patch.object(ai, "make_quiz", fake_quiz):
            d = self.c.post(f"/api/sets/{sid}/quiz", headers=self.h, json={"topic_ids": [ts[0]["id"], ts[2]["id"]], "n": 10, "level": "hard"}).get_json()
        self.assertEqual(len(d["questions"]), 10)
        self.assertEqual({t for t, _ in seen}, {"Cell basics", "Cell division"})
        self.assertEqual({l for _, l in seen}, {"hard"})
        self.assertTrue({q["topic"] for q in d["questions"]} <= {"Cell basics", "Cell division"})
        with mock.patch.object(ai, "make_quiz", fake_quiz):
            odd = self.c.post(f"/api/sets/{sid}/quiz", headers=self.h, json={"n": 6, "level": "banana"}).get_json()
        self.assertEqual(odd["level"], "mixed")

    def test_stats_show_activity_weak_topics_and_streaks(self):
        sid = self.make_set()
        tid = self.topics(sid)[1]["id"]
        self.c.post(f"/api/topics/{tid}/event", headers=self.h, json={"kind": "quiz", "score": 2, "total": 6})
        d = self.c.get("/api/stats", headers=self.h).get_json()
        self.assertEqual(len(d["days"]), 119)
        self.assertGreaterEqual(d["days"][-1]["n"], 1)                                # today's activity counted
        self.assertEqual([w["topic"] for w in d["weak"]], ["Organelles"])
        self.assertEqual((d["totals"]["sets"], d["totals"]["topics"], d["totals"]["quizzes"]), (1, 3, 1))
        self.assertGreaterEqual(d["best_streak"], 1)

    def test_card_editor(self):
        sid, tid = self.first_topic()
        bad = self.c.post(f"/api/topics/{tid}/cards/add", headers=self.h, json={"front": "only a question"})
        self.assertEqual(bad.status_code, 400)
        card = self.c.post(f"/api/topics/{tid}/cards/add", headers=self.h, json={"front": " What is a cell? ", "back": "The unit of life"}).get_json()["card"]
        self.assertEqual(card["front"], "What is a cell?")
        self.assertEqual(self.c.patch(f"/api/cards/{card['id']}", headers=self.h, json={"front": "Define a cell", "back": "Basic unit of life"}).status_code, 200)
        cards = self.c.get(f"/api/sets/{sid}/cards", headers=self.h).get_json()["cards"]
        self.assertEqual([c["front"] for c in cards], ["Define a cell"])
        with mock.patch.object(ai, "make_cards", return_value=[{"front": "Define a cell", "back": "dup"}, {"front": "What is DNA?", "back": "Genes"}]) as more:
            new = self.c.post(f"/api/topics/{tid}/cards/more", headers=self.h).get_json()["cards"]
        self.assertEqual([c["front"] for c in new], ["What is DNA?"])                  # the duplicate was dropped
        self.assertEqual(more.call_args.kwargs["avoid"], ["Define a cell"])
        other = self.c.post("/api/account", json={}).get_json()["key"]
        self.assertEqual(self.c.delete(f"/api/cards/{card['id']}", headers={"Authorization": f"Bearer {other}"}).status_code, 404)   # not theirs
        self.assertEqual(self.c.delete(f"/api/cards/{card['id']}", headers=self.h).status_code, 200)

    def test_tutor_style_reaches_the_ai_only_when_valid(self):
        sid = self.make_set()
        seen = []

        def fake_chat(mode, context, messages, style=""):
            seen.append(style)
            yield "ok"
        with mock.patch.object(ai, "chat", fake_chat):
            for style in ("eli5", "nonsense"):
                self.c.post("/api/chat", headers=self.h, json={"set_id": sid, "style": style, "messages": [{"role": "user", "content": "hi"}]}).get_data()
        self.assertEqual(seen, ["eli5", ""])

    def test_each_free_ai_job_counts_toward_the_cap(self):
        tid = self.topics(self.make_set())[0]["id"]
        with mock.patch.object(ai, "make_swipes", return_value=SWIPES):
            self.c.get(f"/api/topics/{tid}/swipes", headers=self.h)
        uid = server.db.user_for(self.key)["id"]
        self.assertGreaterEqual(server.db.one(server.select(server.db.users).where(server.db.users.c.id == uid))["trial_calls"], 2)


class Access(Base):
    """The free minute, premium links, and the owner's admin tools."""

    def uid(self):
        return server.db.user_for(self.key)["id"]

    def set_user(self, **vals):
        if "trial_used" in vals or "trial_calls" in vals:                               # the free time is per day: mark which day it belongs to
            vals.setdefault("trial_day", server.dt.datetime.now(server.dt.timezone.utc).date().isoformat())
        server.db.run(server.update(server.db.users).where(server.db.users.c.id == self.uid()).values(**vals))

    def admin(self, method, path, **kw):
        with mock.patch.dict(os.environ, {"ADMIN_KEY": "a-long-private-passphrase"}):
            return getattr(self.c, method)(path, headers={"X-Admin-Key": "a-long-private-passphrase"}, **kw)

    def make_link(self, **body):
        d = self.admin("post", "/api/admin/invites", json={"label": "Sara", "uses_max": 2, **body}).get_json()["invite"]
        return d["code"], d["id"]

    def test_free_minute_counts_down_then_locks_everything_but_the_explanation(self):
        d = self.c.get("/api/me", headers=self.h).get_json()
        self.assertEqual((d["premium"], d["locked"], d["trial_total"]), (False, False, server.TRIAL_SECONDS))
        self.assertGreater(d["trial_left"], server.TRIAL_SECONDS - 5)
        for _ in range(3):                                                          # free browsing never runs the clock
            self.assertEqual(self.c.get("/api/sets", headers=self.h).status_code, 200)
        self.assertGreater(self.c.get("/api/me", headers=self.h).get_json()["trial_left"], server.TRIAL_SECONDS - 5)
        self.c.post("/api/beat", headers=self.h, json={"tool": False})
        self.assertGreater(self.c.get("/api/me", headers=self.h).get_json()["trial_left"], server.TRIAL_SECONDS - 5)
        self.c.post("/api/beat", headers=self.h, json={"tool": True})               # but being inside a premium tool does
        self.assertLessEqual(self.c.get("/api/me", headers=self.h).get_json()["trial_left"], server.TRIAL_SECONDS - 1)
        self.set_user(trial_used=server.TRIAL_SECONDS - 0.1, trial_at=server.time.time() - 3)
        r = self.c.post("/api/sets", headers=self.h, data={"text": TEXT})
        self.assertEqual(r.status_code, 402)
        self.assertEqual((r.get_json()["locked"], r.get_json()["contact"]), (True, "Awsaa.libdeh@gmail.com"))
        me = self.c.get("/api/me", headers=self.h).get_json()                      # the page can still ask what's wrong
        self.assertEqual((me["locked"], me["trial_left"]), (True, 0))
        self.assertTrue(self.c.post("/api/beat", headers=self.h).get_json()["locked"])
        self.assertEqual(self.c.post("/api/sets", headers=self.h, data={"text": TEXT}).status_code, 402)

    def test_free_time_comes_back_the_next_day(self):
        self.set_user(trial_used=server.TRIAL_SECONDS, trial_calls=server.TRIAL_CALLS + 1)
        self.assertTrue(self.c.get("/api/me", headers=self.h).get_json()["locked"])
        self.set_user(trial_day="2000-01-01")                                             # yesterday, as far as the server is concerned
        d = self.c.get("/api/me", headers=self.h).get_json()
        self.assertEqual((d["locked"], d["trial_left"]), (False, server.TRIAL_SECONDS))

    def test_only_a_few_ai_jobs_fit_in_the_free_minute(self):
        self.set_user(trial_calls=server.TRIAL_CALLS)
        with mock.patch.object(ai, "make_plan", fake_plan):
            r = self.c.post("/api/sets", headers=self.h, data={"text": TEXT})
        self.assertEqual(r.status_code, 402)                                       # the 9th AI job is refused
        self.assertEqual(self.c.post("/api/sets", headers=self.h, data={"text": TEXT}).status_code, 402)

    def test_a_premium_link_unlocks_and_revoking_it_locks_again(self):
        self.set_user(trial_used=server.TRIAL_SECONDS)
        self.assertEqual(self.c.post("/api/sets", headers=self.h, data={"text": TEXT}).status_code, 402)
        code, iid = self.make_link()
        r = self.c.post("/api/join", headers=self.h, json={"code": code})
        self.assertEqual((r.status_code, r.get_json()["premium"]), (200, True))
        me = self.c.get("/api/me", headers=self.h).get_json()
        self.assertEqual((me["premium"], me["locked"]), (True, False))
        self.assertEqual(self.c.get("/api/sets", headers=self.h).status_code, 200)
        self.admin("post", f"/api/admin/invites/{iid}/revoke", json={"revoked": True})
        self.assertEqual(self.c.post("/api/sets", headers=self.h, data={"text": TEXT}).status_code, 402)  # switching the link off takes it back
        self.admin("post", f"/api/admin/invites/{iid}/revoke", json={"revoked": False})
        self.assertEqual(self.c.get("/api/sets", headers=self.h).status_code, 200)

    def test_link_device_limit_and_unknown_codes(self):
        code, _ = self.make_link()
        keys = [self.c.post("/api/join", json={"code": code}).get_json()["key"] for _ in range(2)]   # a new device needs no account first
        self.assertTrue(all(keys))
        for k in keys:
            self.assertTrue(self.c.get("/api/me", headers={"Authorization": f"Bearer {k}"}).get_json()["premium"])
        third = self.c.post("/api/join", json={"code": code})
        self.assertEqual(third.status_code, 400)
        self.assertIn("Awsaa.libdeh@gmail.com", third.get_json()["error"])
        self.assertEqual(self.c.post("/api/join", json={"code": "nope"}).status_code, 400)
        again = self.c.post("/api/join", headers={"Authorization": f"Bearer {keys[0]}"}, json={"code": code})
        self.assertEqual(again.status_code, 200)                                    # the same device re-opening it costs nothing

    def test_premium_can_end_after_some_days(self):
        code, _ = self.make_link(days=30)
        self.c.post("/api/join", headers=self.h, json={"code": code})
        self.assertAlmostEqual(self.c.get("/api/me", headers=self.h).get_json()["premium_until"], server.time.time() + 30 * 86400, delta=60)
        self.set_user(premium_until=server.time.time() - 10, trial_used=server.TRIAL_SECONDS)
        self.assertEqual(self.c.post("/api/sets", headers=self.h, data={"text": TEXT}).status_code, 402)

    def test_admin_needs_the_secret_key(self):
        self.assertEqual(self.c.get("/api/admin/overview").status_code, 403)                  # no key configured: off
        with mock.patch.dict(os.environ, {"ADMIN_KEY": "a-long-private-passphrase"}):
            self.assertEqual(self.c.get("/api/admin/overview", headers={"X-Admin-Key": "guess"}).status_code, 403)
            self.assertEqual(self.c.get("/api/admin/overview", headers={"X-Admin-Key": "a-long-private-passphrase"}).status_code, 200)
        with mock.patch.dict(os.environ, {"ADMIN_KEY": "short"}):                              # a weak key is as good as none
            self.assertEqual(self.c.get("/api/admin/overview", headers={"X-Admin-Key": "short"}).status_code, 403)
        code, iid = self.make_link(label="Omar")
        ov = self.admin("get", "/api/admin/overview").get_json()
        mine = [i for i in ov["invites"] if i["id"] == iid]
        self.assertEqual([i["label"] for i in mine], ["Omar"])
        self.assertTrue(mine[0]["link"].endswith(f"/join/{code}"))
        self.assertGreaterEqual(ov["accounts"], 1)
        self.admin("delete", f"/api/admin/invites/{iid}")
        self.assertNotIn(iid, [i["id"] for i in self.admin("get", "/api/admin/overview").get_json()["invites"]])

    def test_join_page_and_admin_page_are_served(self):
        page = self.c.get("/join/abc123")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'window.JOIN = "abc123"', page.data)
        self.assertEqual(page.headers["Referrer-Policy"], "no-referrer")                       # the code never leaks to other sites
        self.assertEqual(self.c.get("/admin").status_code, 200)
        self.assertIn(b"window.JOIN = null", self.c.get("/").data)

    def test_only_a_few_free_accounts_per_day_from_one_address(self):
        for _ in range(server.FREE_ACCOUNTS_PER_DAY - 1):                           # setUp already made one
            self.assertEqual(self.c.post("/api/account", json={}).status_code, 200)
        r = self.c.post("/api/account", json={})
        self.assertEqual(r.status_code, 429)
        self.assertTrue(r.get_json()["locked"])


class OwnerSwitches(Base):
    """Ban, time-out, messages, the kill switch and personal links."""
    uid, admin, make_link = Access.uid, Access.admin, Access.make_link

    def me(self):
        return self.c.get("/api/me", headers=self.h).get_json()

    def test_the_owner_is_never_rate_limited_but_wrong_keys_are(self):
        with mock.patch.dict(os.environ, {"ADMIN_KEY": "a-long-private-passphrase"}):
            for _ in range(90):
                self.assertEqual(self.c.get("/api/admin/overview", headers={"X-Admin-Key": "a-long-private-passphrase"}).status_code, 200)
            codes = {self.c.get("/api/admin/overview", headers={"X-Admin-Key": "guess"}).status_code for _ in range(40)}
            self.assertEqual(codes, {403, 429})                                          # guessing gets shut out
            self.assertEqual(self.c.get("/api/admin/overview", headers={"X-Admin-Key": "a-long-private-passphrase"}).status_code, 200)

    def test_ban_blocks_everything_but_explains_itself(self):
        uid = self.uid()
        self.admin("post", f"/api/admin/users/{uid}/ban", json={"banned": True, "reason": "Spamming"})
        r = self.c.get("/api/sets", headers=self.h)
        self.assertEqual((r.status_code, r.get_json()["banned"], r.get_json()["reason"]), (403, True, "Spamming"))
        self.assertTrue(self.me()["banned"])                                           # /api/me still answers, so the page can explain
        self.admin("post", f"/api/admin/users/{uid}/ban", json={"banned": False})
        self.assertEqual(self.c.get("/api/sets", headers=self.h).status_code, 200)

    def test_timeout_ends_by_itself(self):
        uid = self.uid()
        self.admin("post", f"/api/admin/users/{uid}/timeout", json={"minutes": 30, "reason": "Cool off"})
        r = self.c.get("/api/sets", headers=self.h)
        self.assertEqual((r.status_code, r.get_json()["timeout"]), (403, True))
        self.assertGreater(self.me()["timeout_until"], time.time())
        server.db.run(server.update(server.db.users).where(server.db.users.c.id == uid).values(timeout_until=time.time() - 5))
        self.assertEqual(self.c.get("/api/sets", headers=self.h).status_code, 200)
        self.assertEqual(self.me()["timeout_until"], 0)

    def test_messages_arrive_once_and_broadcast_reaches_everyone(self):
        uid = self.uid()
        self.assertEqual(self.admin("post", f"/api/admin/users/{uid}/message", json={"body": "  "}).status_code, 400)
        self.admin("post", f"/api/admin/users/{uid}/message", json={"body": "Hi Aws, thanks for testing!"})
        self.assertEqual([m["body"] for m in self.me()["inbox"]], ["Hi Aws, thanks for testing!"])
        self.c.post("/api/messages/read", headers=self.h)
        self.assertEqual(self.me()["inbox"], [])
        sent = self.admin("post", "/api/admin/broadcast", json={"body": "New feature!"}).get_json()["sent"]
        self.assertGreaterEqual(sent, 1)
        self.assertEqual([m["body"] for m in self.me()["inbox"]], ["New feature!"])

    def test_kill_switch_stops_the_app_but_not_the_owner(self):
        self.admin("post", "/api/admin/settings", json={"maintenance": True, "message": "Back at 6pm"})
        try:
            r = self.c.get("/api/sets", headers=self.h)
            self.assertEqual((r.status_code, r.get_json()["maintenance"], r.get_json()["error"]), (503, True, "Back at 6pm"))
            self.assertTrue(self.c.get("/api/status").get_json()["maintenance"])
            self.assertEqual(self.c.post("/api/account", json={}).status_code, 503)    # no new accounts either
            self.assertTrue(self.me()["maintenance"])
            self.assertEqual(self.admin("get", "/api/admin/overview").status_code, 200)
        finally:
            self.admin("post", "/api/admin/settings", json={"maintenance": False})
        self.assertEqual(self.c.get("/api/sets", headers=self.h).status_code, 200)

    def test_personal_link_carries_the_note(self):
        code, _ = self.make_link(label="Sara", message="Hey Sara, this one is just for you!")
        d = self.c.post("/api/join", json={"code": code}).get_json()
        self.assertEqual((d["message"], d["for_name"]), ("Hey Sara, this one is just for you!", "Sara"))

    def test_owner_can_grant_reset_and_delete(self):
        uid = self.uid()
        self.admin("post", f"/api/admin/users/{uid}/premium", json={"on": True, "days": 7})
        self.assertTrue(self.me()["premium"])
        self.admin("post", f"/api/admin/users/{uid}/premium", json={"on": False})
        self.assertFalse(self.me()["premium"])
        users = self.admin("get", "/api/admin/users").get_json()["users"]
        self.assertIn(uid, [u["id"] for u in users])
        self.assertEqual(self.admin("delete", f"/api/admin/users/{uid}").status_code, 200)
        self.assertEqual(self.c.get("/api/me", headers=self.h).status_code, 401)


class Universe(Base):
    """One account in every app: the other apps ask /api/uni/guard what each visitor should see."""
    uid, admin = Access.uid, Access.admin

    def guard(self, app_id="alibi", key=True, dev="d1", since=0):
        h = {"Authorization": f"Bearer {self.key}"} if key else {}
        return self.c.get(f"/api/uni/guard?app={app_id}&dev={dev}&since={since}", headers=h).get_json()

    def test_an_unlinked_visitor_gets_a_clean_answer_and_cors(self):
        r = self.c.get("/api/uni/guard?app=alibi&dev=x")
        self.assertEqual(r.headers["Access-Control-Allow-Origin"], "*")
        j = r.get_json()
        self.assertEqual((j["off"], j["banned"], j["linked"], j["notes"]), (False, False, False, []))
        self.assertEqual(self.c.open("/api/uni/guard", method="OPTIONS").status_code, 204)

    def test_linking_with_a_pairing_code_gives_the_same_account(self):
        code = self.c.post("/api/pair", headers=self.h).get_json()["code"]
        got = self.c.post("/api/uni/link", json={"code": code}).get_json()
        self.assertTrue(got["key"])
        self.assertEqual(got["name"], "Aws")
        self.assertEqual(self.c.post("/api/uni/link", json={"code": code}).status_code, 400)       # a code works once
        j = self.c.get("/api/uni/guard?app=scout", headers={"Authorization": f"Bearer {got['key']}"}).get_json()
        self.assertEqual((j["linked"], j["name"]), (True, "Aws"))

    def test_a_ban_made_in_cramly_follows_the_account_into_every_app(self):
        uid = self.uid()
        self.admin("post", f"/api/admin/users/{uid}/ban", json={"banned": True, "reason": "Trolling"})
        for app_id in ("alibi", "scout", "platter"):
            j = self.guard(app_id)
            self.assertEqual((j["banned"], j["reason"]), (True, "Trolling"), app_id)
        self.admin("post", f"/api/admin/users/{uid}/ban", json={"banned": False})
        self.assertFalse(self.guard()["banned"])

    def test_a_timeout_follows_the_account_too(self):
        self.admin("post", f"/api/admin/users/{self.uid()}/timeout", json={"minutes": 20, "reason": "Break"})
        self.assertGreater(self.guard("homebase")["timeout_until"], time.time())

    def test_switching_one_app_off_leaves_the_others_alone(self):
        self.admin("post", "/api/admin/uni/app", json={"app": "alibi", "off": True, "msg": "Fixing a bug"})
        self.assertEqual((self.guard("alibi")["off"], self.guard("alibi")["message"]), (True, "Fixing a bug"))
        self.assertFalse(self.guard("scout")["off"])
        self.admin("post", "/api/admin/uni/app", json={"app": "alibi", "off": False})
        self.assertFalse(self.guard("alibi")["off"])

    def test_switching_everything_off(self):
        self.admin("post", "/api/admin/uni/app", json={"app": "*", "off": True, "msg": "Back soon"})
        self.assertTrue(all(self.guard(a)["off"] for a in ("alibi", "scout", "cramly")))
        self.assertEqual(self.c.get("/api/sets", headers=self.h).status_code, 503)               # Cramly itself obeys it
        self.admin("post", "/api/admin/uni/app", json={"app": "*", "off": False})
        self.assertEqual(self.c.get("/api/sets", headers=self.h).status_code, 200)

    def test_messages_and_trolls_reach_only_the_chosen_app_and_only_once(self):
        self.admin("post", "/api/admin/uni/note", json={"app": "alibi", "kind": "message", "body": "Hello players"})
        self.admin("post", "/api/admin/uni/note", json={"app": "*", "kind": "troll", "effect": "flip"})
        first = self.guard("alibi")
        self.assertEqual([n["kind"] for n in first["notes"]], ["message", "troll"])
        self.assertEqual([n["kind"] for n in self.guard("scout")["notes"]], ["troll"])
        self.assertEqual(self.guard("alibi", since=first["last"])["notes"], [])                  # already seen
        self.assertEqual(self.admin("post", "/api/admin/uni/note", json={"app": "x", "kind": "message", "body": "a"}).status_code, 404)

    def test_a_visitor_without_an_account_can_be_blocked_by_device(self):
        self.admin("post", "/api/admin/uni/device", json={"dev": "abc123", "reason": "Spam"})
        j = self.guard(key=False, dev="abc123")
        self.assertEqual((j["banned"], j["reason"]), (True, "Spam"))
        self.assertFalse(self.guard(key=False, dev="someone-else")["banned"])
        self.admin("post", "/api/admin/uni/device", json={"dev": "abc123", "clear": True})
        self.assertFalse(self.guard(key=False, dev="abc123")["banned"])

    def test_the_inbox_shows_in_other_apps_until_it_is_read(self):
        self.admin("post", f"/api/admin/users/{self.uid()}/message", json={"body": "Hi from the owner"})
        self.assertEqual([m["body"] for m in self.guard("alibi")["inbox"]], ["Hi from the owner"])
        self.c.post("/api/uni/ack", headers=self.h)
        self.assertEqual(self.guard("alibi")["inbox"], [])

    def test_other_apps_can_ask_whether_a_key_is_the_owners(self):
        self.assertEqual(self.c.get("/api/uni/verify").status_code, 403)
        self.assertEqual(self.c.get("/api/uni/verify", headers={"X-Admin-Key": "wrong-wrong-wrong"}).status_code, 403)
        self.assertEqual(self.admin("get", "/api/uni/verify").status_code, 200)

    def test_owner_routes_need_the_key(self):
        self.assertEqual(self.c.get("/api/admin/uni").status_code, 403)


class BigUploads(Base):
    """Files up to 1 GB are streamed to disk and read from there."""

    def test_big_text_file_is_read_without_loading_it_all(self):
        import tempfile
        with tempfile.NamedTemporaryFile(delete=False, suffix=".txt") as f:
            f.write(b"Mitochondria make ATP for the cell. " * 600_000)       # about 22 MB
            path = f.name
        try:
            kind, text = ingest.read_path("big.txt", path)
        finally:
            os.remove(path)
        self.assertEqual(kind, "text")
        self.assertLessEqual(len(text), ingest.MAX_TEXT)

    def test_limits_are_one_gigabyte_and_twenty_for_photos(self):
        self.assertEqual(ingest.MAX_FILE, 1024 ** 3)
        import tempfile
        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as f:
            f.write(b"x" * (21 * 1024 * 1024))
            path = f.name
        try:
            with self.assertRaises(ingest.ai.AIError) as e:
                ingest.read_path("photo.png", path)
            self.assertIn("20 MB", str(e.exception))
        finally:
            os.remove(path)

    def test_oversized_request_gets_a_clear_message(self):
        with mock.patch.dict(server.app.config, {"MAX_CONTENT_LENGTH": 1000}):
            r = self.c.post("/api/sets", headers=self.h, data={"files": (io.BytesIO(b"a" * 5000), "n.txt")})
        self.assertEqual(r.status_code, 413)
        self.assertIn("1 GB", r.get_json()["error"])

    def test_uploaded_files_do_not_pile_up_in_temp(self):
        import tempfile
        before = {n for n in os.listdir(tempfile.gettempdir()) if n.startswith("cramly_up_")}
        with mock.patch.object(ai, "make_plan", return_value=PLAN), mock.patch.object(ai, "make_notes", return_value="n"):
            self.c.post("/api/sets", headers=self.h, data={"files": (io.BytesIO(("Cells are the basic unit of life. " * 40).encode()), "n.txt")}).get_data()
        after = {n for n in os.listdir(tempfile.gettempdir()) if n.startswith("cramly_up_")}
        self.assertEqual(after - before, set())


class Compression(Base):
    """Squeezing text before the AI reads it, and accepting gzip uploads."""

    def test_squeeze_drops_headers_page_numbers_and_repeats_but_keeps_ideas(self):
        nl = chr(10)
        rows = ["ACME Corp confidential"]
        for i in range(1, 12):
            rows += [f"Page {i}", "ACME Corp confidential", f"Unique idea number {i} about cells and energy flow."]
        out = ingest.squeeze(nl.join(rows))
        self.assertEqual((out.count("ACME"), out.count("Page")), (1, 0))
        self.assertEqual(out.count("Unique idea"), 11)

    def test_gzip_upload_is_unpacked_and_read(self):
        import gzip
        text = ("Photosynthesis turns light into sugar. " * 30).encode()
        kind, got = ingest.read_file("notes.txt.gz", gzip.compress(text))
        self.assertEqual(kind, "text")
        self.assertIn("Photosynthesis turns light", got)

    def test_gzip_bomb_is_refused_and_bad_gzip_is_friendly(self):
        import gzip
        with mock.patch.object(ingest, "MAX_GUNZIP", 1000):
            with self.assertRaises(ingest.ai.AIError) as e:
                ingest.read_file("big.txt.gz", gzip.compress(b"a" * 5000))
        self.assertIn("unpacks to more", str(e.exception))
        with self.assertRaises(ingest.ai.AIError):
            ingest.read_file("broken.txt.gz", b"not gzip at all")

    def test_gz_upload_through_the_api_builds_a_set(self):
        import gzip
        body = gzip.compress(("Cells are the basic unit of life. " * 40).encode())
        with mock.patch.object(ai, "make_plan", return_value=PLAN), mock.patch.object(ai, "make_notes", return_value="n"):
            r = self.c.post("/api/sets", headers=self.h, data={"files": (io.BytesIO(body), "notes.txt.gz")}).get_data(as_text=True)
        self.assertIn('"type": "set"', r)


class Tiers(Base):
    """Free, Premium and Premium Plus, and the Plus tools."""
    uid, admin, make_link, set_user = Access.uid, Access.admin, Access.make_link, Access.set_user

    def plan(self):
        return self.c.get("/api/me", headers=self.h).get_json()

    def join(self, **link):
        code, _ = self.make_link(**link)
        self.c.post("/api/join", headers=self.h, json={"code": code})

    def test_plans_and_limits(self):
        d = self.plan()
        self.assertEqual((d["plan"], d["limits"]["sets"], d["limits"]["file_mb"]), ("free", 2, 10))
        self.join()
        d = self.plan()
        self.assertEqual((d["plan"], d["limits"]["sets"], d["limits"]["file_mb"]), ("premium", 30, 200))
        self.join(plan="plus")
        d = self.plan()
        self.assertEqual((d["plan"], d["plus"], d["limits"]["sets"], d["limits"]["file_mb"]), ("plus", True, 100, 1024))

    def test_plus_tools_free_minute_then_plus_only(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        with mock.patch.object(ai, "make_mixups", return_value=[{"a": "mitosis", "b": "meiosis", "difference": "d", "tip": "t"}]):
            self.assertEqual(self.c.get(f"/api/sets/{sid}/mixups", headers=self.h).status_code, 200)        # inside the free minute
            self.set_user(trial_used=server.TRIAL_SECONDS)
            r = self.c.get(f"/api/sets/{sid}/mixups?fresh=1", headers=self.h)
            self.assertEqual((r.status_code, r.get_json()["plus"]), (402, True))                            # minute is over
            self.join()                                                                                      # premium, not plus
            r = self.c.get(f"/api/sets/{sid}/mixups", headers=self.h)
            self.assertEqual((r.status_code, r.get_json()["error"]), (402, "This tool is part of Premium Plus."))
            self.join(plan="plus")
            self.assertEqual(self.c.get(f"/api/sets/{sid}/mixups", headers=self.h).status_code, 200)
        self.assertEqual(tid > 0, True)

    def test_free_keeps_reading_and_the_basics_after_the_minute(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        self.set_user(trial_used=server.TRIAL_SECONDS, trial_at=server.time.time() - 3)
        for path in ("/api/sets", f"/api/sets/{sid}", "/api/stats", f"/api/sets/{sid}/cards"):
            self.assertEqual(self.c.get(path, headers=self.h).status_code, 200, path)
        for path in (f"/api/sets/{sid}/mindmap", "/api/stats?full=1", f"/api/topics/{tid}/mynotes"):       # Premium now
            self.assertEqual(self.c.get(path, headers=self.h).status_code, 402, path)
        self.join()
        for path in (f"/api/sets/{sid}/mindmap", "/api/stats?full=1", f"/api/topics/{tid}/mynotes"):
            self.assertEqual(self.c.get(path, headers=self.h).status_code, 200, path)

    def test_glossary_is_cached_for_premium(self):
        sid = self.make_set()
        self.join()
        terms = [{"term": "Cell", "definition": "Basic unit of life."}]
        with mock.patch.object(ai, "make_glossary", return_value=terms) as made:
            for _ in range(2):
                d = self.c.get(f"/api/sets/{sid}/glossary", headers=self.h).get_json()
        self.assertEqual((d["terms"], made.call_count), (terms, 1))

    def test_grader_and_prompt(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        self.join(plan="plus")
        self.assertEqual(self.c.post(f"/api/topics/{tid}/grade", headers=self.h, json={"prompt": "Q?", "answer": "short"}).status_code, 400)
        graded = {"score": 72, "verdict": "Good start", "strengths": ["a"], "fixes": ["b"], "improved": "better"}
        with mock.patch.object(ai, "grade_answer", return_value=graded), mock.patch.object(ai, "make_prompt", return_value="Explain the cell."):
            self.assertEqual(self.c.post(f"/api/topics/{tid}/prompt", headers=self.h).get_json()["prompt"], "Explain the cell.")
            r = self.c.post(f"/api/topics/{tid}/grade", headers=self.h, json={"prompt": "Explain the cell.", "answer": "A cell is the basic unit of life and has parts."})
        self.assertEqual(r.get_json()["score"], 72)

    def test_practice_lab_and_boost(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        self.join(plan="plus")
        lab = {"cloze": [{"text": "The ____ makes ATP.", "answer": "mitochondrion", "hint": "powerhouse"}], "tf": [{"statement": "x", "answer": True, "why": "y"}]}
        boost = {"analogies": ["a"], "mnemonics": ["m"], "example": "e", "common_mistake": "c"}
        with mock.patch.object(ai, "make_lab", return_value=lab), mock.patch.object(ai, "make_boost", return_value=boost):
            self.assertEqual(self.c.get(f"/api/topics/{tid}/lab", headers=self.h).get_json()["lab"], lab)
            self.assertEqual(self.c.get(f"/api/topics/{tid}/boost", headers=self.h).get_json()["boost"], boost)

    def test_sharing_is_plus_only_and_import_copies_without_progress(self):
        sid = self.make_set()
        self.assertEqual(self.c.post(f"/api/sets/{sid}/share", headers=self.h, json={"on": True}).status_code, 200)   # free minute
        self.join()
        self.assertEqual(self.c.post(f"/api/sets/{sid}/share", headers=self.h, json={"on": True}).status_code, 402)   # premium is not enough
        self.join(plan="plus")
        token = self.c.post(f"/api/sets/{sid}/share", headers=self.h, json={"on": True}).get_json()["token"]
        self.assertEqual(self.c.get(f"/api/shared/{token}").get_json()["title"], "Cells")
        friend = self.c.post("/api/account", json={"name": "Friend"}).get_json()["key"]
        fh = {"Authorization": f"Bearer {friend}"}
        new_id = self.c.post(f"/api/shared/{token}/import", headers=fh).get_json()["id"]
        copy = self.c.get(f"/api/sets/{new_id}", headers=fh).get_json()
        self.assertEqual(len(copy["topics"]), 3)
        self.assertTrue(all(t["status"] == 0 for t in copy["topics"]))
        self.c.post(f"/api/sets/{sid}/share", headers=self.h, json={"on": False})
        self.assertEqual(self.c.get(f"/api/shared/{token}").status_code, 404)

    def test_studio_voice_is_plus_only_and_validated(self):
        with mock.patch.object(ai, "speech", return_value=b"ID3fakemp3") as made:
            self.assertEqual(self.c.post("/api/tts", headers=self.h, json={"text": "Hello there"}).status_code, 402)      # free
            self.join()
            self.assertEqual(self.c.post("/api/tts", headers=self.h, json={"text": "Hello there"}).status_code, 402)      # premium
            self.join(plan="plus")
            self.assertEqual(self.c.post("/api/tts", headers=self.h, json={"text": " "}).status_code, 400)
            r = self.c.post("/api/tts", headers=self.h, json={"text": "Hello " * 300, "voice": "nova"})
            self.assertEqual((r.status_code, r.mimetype, r.data), (200, "audio/mpeg", b"ID3fakemp3"))
            self.assertLessEqual(len(made.call_args.args[0]), 600)
            self.assertEqual(made.call_args.args[1], "nova")

    def test_free_set_cap_and_file_cap(self):
        for _ in range(2):
            self.make_set()
        with mock.patch.object(ai, "make_plan", return_value=PLAN):
            r = self.c.post("/api/sets", headers=self.h, data={"text": "Cells are the basic unit of life. " * 20})
        self.assertEqual((r.status_code, r.get_json()["plan_limit"]), (400, True))
        self.join()
        big = io.BytesIO(b"x" * (201 * 1024 * 1024))
        r = self.c.post("/api/sets", headers=self.h, data={"files": (big, "huge.txt")})
        self.assertEqual((r.status_code, r.get_json()["plan_limit"]), (413, True))
        self.assertIn("Premium Plus", r.get_json()["error"])


class PlanRequests(Base):
    """Asking for a plan inside the app, and the owner deciding."""
    uid, admin = Access.uid, Access.admin

    def me(self):
        return self.c.get("/api/me", headers=self.h).get_json()

    def test_ask_then_owner_gives_it_free(self):
        uid = self.uid()
        r = self.c.post("/api/request", headers=self.h, json={"plan": "plus", "message": "I am a student with an exam", "name": "Sam"})
        self.assertEqual((r.status_code, r.get_json()["status"]), (200, "pending"))
        self.assertEqual(self.me()["request"]["status"], "pending")
        items = self.admin("get", "/api/admin/requests").get_json()
        mine = next(i for i in items["requests"] if i["user_id"] == uid)
        self.assertEqual((mine["plan"], mine["message"], mine["status"]), ("plus", "I am a student with an exam", "pending"))
        self.assertGreaterEqual(items["pending"], 1)
        self.assertGreaterEqual(self.admin("get", "/api/admin/overview").get_json()["pending_requests"], 1)
        done = self.admin("post", f"/api/admin/requests/{mine['id']}/decide", json={"decision": "free", "note": "Good luck!"})
        self.assertEqual(done.get_json()["status"], "free")
        d = self.me()
        self.assertEqual((d["plan"], d["request"]["status"]), ("plus", "free"))
        self.assertIn("for free", d["inbox"][-1]["body"])
        self.assertIn("Good luck!", d["inbox"][-1]["body"])

    def test_owner_can_ask_for_payment_then_mark_it_paid(self):
        self.c.post("/api/request", headers=self.h, json={"plan": "premium", "message": "please"})
        rid = self.admin("get", "/api/admin/requests").get_json()["requests"][0]["id"]
        self.admin("post", f"/api/admin/requests/{rid}/decide", json={"decision": "pay", "price": "$2.99", "note": "Send it to my PayPal."})
        d = self.me()
        self.assertEqual((d["plan"], d["request"]["status"], d["request"]["price"]), ("free", "pay", "$2.99"))
        self.assertIn("$2.99", d["inbox"][-1]["body"])
        self.admin("post", f"/api/admin/requests/{rid}/decide", json={"decision": "paid", "days": 30})
        d = self.me()
        self.assertEqual((d["plan"], d["request"]["status"]), ("premium", "paid"))
        self.assertGreater(d["premium_until"], server.time.time())

    def test_decline_and_validation(self):
        self.c.post("/api/request", headers=self.h, json={"message": "hi"})
        rid = self.admin("get", "/api/admin/requests").get_json()["requests"][0]["id"]
        self.assertEqual(self.admin("post", f"/api/admin/requests/{rid}/decide", json={"decision": "maybe"}).status_code, 400)
        self.assertEqual(self.admin("post", "/api/admin/requests/99999/decide", json={"decision": "free"}).status_code, 404)
        self.admin("post", f"/api/admin/requests/{rid}/decide", json={"decision": "declined"})
        self.assertEqual(self.me()["request"]["status"], "declined")
        self.assertEqual(self.me()["plan"], "free")

    def test_asking_again_updates_the_same_request_and_is_rate_limited(self):
        for i in range(2):
            self.c.post("/api/request", headers=self.h, json={"message": f"try {i}"})
        mine = [r for r in self.admin("get", "/api/admin/requests").get_json()["requests"] if r["user_id"] == self.uid()]
        self.assertEqual((len(mine), mine[0]["message"]), (1, "try 1"))
        for _ in range(4):                                                          # four separate requests in a day are fine
            rid = [r for r in self.admin("get", "/api/admin/requests").get_json()["requests"] if r["user_id"] == self.uid()][0]["id"]
            self.admin("post", f"/api/admin/requests/{rid}/decide", json={"decision": "declined"})
            self.assertEqual(self.c.post("/api/request", headers=self.h, json={"message": "more"}).status_code in (200, 429), True)
        rid = [r for r in self.admin("get", "/api/admin/requests").get_json()["requests"] if r["user_id"] == self.uid()][0]["id"]
        self.admin("post", f"/api/admin/requests/{rid}/decide", json={"decision": "declined"})
        self.assertEqual(self.c.post("/api/request", headers=self.h, json={"message": "again"}).status_code, 429)

    def test_questions_bugs_and_ideas_get_replies(self):
        self.assertEqual(self.c.post("/api/request", headers=self.h, json={"kind": "bug", "message": "x"}).status_code, 400)
        self.c.post("/api/request", headers=self.h, json={"kind": "bug", "message": "The match game freezes on my phone."})
        self.c.post("/api/request", headers=self.h, json={"kind": "premium", "plan": "premium", "message": "please"})
        rows = [r for r in self.admin("get", "/api/admin/requests").get_json()["requests"] if r["user_id"] == self.uid()]
        self.assertEqual({r["kind"] for r in rows}, {"bug", "premium"})                       # two separate threads
        bug = next(r for r in rows if r["kind"] == "bug")
        self.assertEqual(self.admin("post", f"/api/admin/requests/{bug['id']}/decide", json={"decision": "reply"}).status_code, 400)
        self.admin("post", f"/api/admin/requests/{bug['id']}/decide", json={"decision": "reply", "note": "Thanks, fixed in the next update."})
        d = self.me()
        self.assertTrue(any("fixed in the next update" in m["body"] for m in d["inbox"]))
        after = {r["id"]: r["status"] for r in self.admin("get", "/api/admin/requests").get_json()["requests"]}
        self.assertEqual(after[bug["id"]], "answered")

    def test_cannot_ask_for_what_you_have(self):
        code, _ = Access.make_link(self, plan="plus")
        self.c.post("/api/join", headers=self.h, json={"code": code})
        self.assertEqual(self.c.post("/api/request", headers=self.h, json={"plan": "plus"}).status_code, 400)


class MoreTools(Base):
    """Summariser, quick cards, outline, card improver, study guide."""
    uid, admin, make_link, set_user = Access.uid, Access.admin, Access.make_link, Access.set_user

    def join(self, **link):
        code, _ = self.make_link(**link)
        self.c.post("/api/join", headers=self.h, json={"code": code})

    def test_summarize_needs_text_and_is_premium(self):
        summary = {"short": "s", "bullets": ["a"], "detailed": "d", "terms": ["t"]}
        with mock.patch.object(ai, "summarize", return_value=summary):
            self.assertEqual(self.c.post("/api/summarize", headers=self.h, json={"text": "too short"}).status_code, 400)
            self.assertEqual(self.c.post("/api/summarize", headers=self.h, json={"text": "Cells are the unit of life. " * 10}).get_json()["short"], "s")   # free minute
            self.set_user(trial_used=server.TRIAL_SECONDS)
            self.assertEqual(self.c.post("/api/summarize", headers=self.h, json={"text": "Cells are the unit of life. " * 10}).status_code, 402)
            self.join()
            self.assertEqual(self.c.post("/api/summarize", headers=self.h, json={"text": "Cells are the unit of life. " * 10}).status_code, 200)

    def test_quick_cards_can_be_saved_into_a_topic_without_duplicates(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        cards = [{"front": "What is a cell?", "back": "The unit of life."}, {"front": "What is DNA?", "back": "Genes."}]
        with mock.patch.object(ai, "quick_cards", return_value=cards):
            text = "Cells are the unit of life. " * 10
            self.assertEqual(self.c.post("/api/quickcards", headers=self.h, json={"text": text}).get_json()["saved"], 0)
            r = self.c.post("/api/quickcards", headers=self.h, json={"text": text, "topic_id": tid, "save": True}).get_json()
            self.assertEqual(r["saved"], 2)
            again = self.c.post("/api/quickcards", headers=self.h, json={"text": text, "topic_id": tid, "save": True}).get_json()
            self.assertEqual(again["saved"], 0)

    def test_outline_improver_and_guide_are_plus(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        self.join()
        self.assertEqual(self.c.get(f"/api/sets/{sid}/guide", headers=self.h).status_code, 402)
        self.join(plan="plus")
        outline = {"thesis": "t", "sections": [{"heading": "h", "points": ["p"]}], "evidence": ["e"], "conclusion": "c"}
        with mock.patch.object(ai, "make_essay_outline", return_value=outline):
            self.assertEqual(self.c.post(f"/api/sets/{sid}/outline", headers=self.h, json={"question": "hi"}).status_code, 400)
            self.assertEqual(self.c.post(f"/api/sets/{sid}/outline", headers=self.h, json={"question": "Discuss how cells make energy."}).get_json()["thesis"], "t")
        cid = self.c.post(f"/api/topics/{tid}/cards/add", headers=self.h, json={"front": "Q?", "back": "A."}).get_json()["card"]["id"]
        better = [{"id": cid, "front": "Clearer Q?", "back": "Crisp A.", "tip": "Think of a factory."}]
        with mock.patch.object(ai, "improve_cards", return_value=better):
            out = self.c.post(f"/api/topics/{tid}/cards/improve", headers=self.h).get_json()["cards"]
        self.assertEqual(out[0]["tip"], "Think of a factory.")
        self.assertEqual(self.c.get(f"/api/sets/{sid}/cards", headers=self.h).get_json()["cards"][0]["front"], "Clearer Q?")
        g = self.c.get(f"/api/sets/{sid}/guide", headers=self.h).get_json()
        self.assertEqual(g["topics"], 3)
        self.assertIn("## 1. Cell basics", g["guide"])


class Tools4(Base):
    """Daily challenge, smart drill, card import and rewrite."""
    uid, admin, make_link, set_user = Access.uid, Access.admin, Access.make_link, Access.set_user

    def join(self, **link):
        code, _ = self.make_link(**link)
        self.c.post("/api/join", headers=self.h, json={"code": code})

    def test_daily_challenge_comes_from_your_own_quizzes_and_is_stable(self):
        self.join()
        self.assertTrue(self.c.get("/api/daily", headers=self.h).get_json()["none"])             # nothing to pick from yet
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        with mock.patch.object(ai, "make_quiz", return_value=[dict(q) for q in QUIZ]):
            self.c.get(f"/api/topics/{tid}/quiz", headers=self.h)
        a = self.c.get("/api/daily", headers=self.h).get_json()
        b = self.c.get("/api/daily", headers=self.h).get_json()
        self.assertEqual((a["q"], a["topic"], a["set_title"]), (b["q"], "Cell basics", "Cells"))
        self.assertEqual(len(a["options"]), 4)

    def test_drill_returns_weak_cards_from_every_set(self):
        self.join()
        for _ in range(2):
            sid = self.make_set()
            tid = self.topics(sid)[0]["id"]
            self.c.post(f"/api/topics/{tid}/cards/add", headers=self.h, json={"front": f"Q{sid}?", "back": "A."})
        d = self.c.get("/api/drill", headers=self.h).get_json()
        self.assertEqual((d["weak"], len(d["cards"])), (2, 2))
        self.assertEqual({c["set_title"] for c in d["cards"]}, {"Cells"})

    def test_bulk_import_understands_tabs_pipes_semicolons_and_skips_duplicates(self):
        self.join()
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        text = chr(10).join(["What is a cell?" + chr(9) + "The unit of life", "What is DNA? | Genes", "What is RNA?;A copy of genes", "no separator here", "What is a cell?" + chr(9) + "dup"])
        r = self.c.post(f"/api/topics/{tid}/cards/bulk", headers=self.h, json={"text": text}).get_json()
        self.assertEqual((r["added"], r["skipped"]), (3, 2))
        self.assertEqual(self.c.post(f"/api/topics/{tid}/cards/bulk", headers=self.h, json={"text": "nothing useful"}).status_code, 400)

    def test_rewrite_is_plus_needs_notes_and_valid_mode(self):
        sid = self.make_set()
        tid = self.topics(sid)[0]["id"]
        self.join()
        self.assertEqual(self.c.post(f"/api/topics/{tid}/rewrite", headers=self.h, json={"mode": "simpler"}).status_code, 402)
        self.join(plan="plus")
        self.assertEqual(self.c.post(f"/api/topics/{tid}/rewrite", headers=self.h, json={"mode": "banana"}).status_code, 400)
        self.assertEqual(self.c.post(f"/api/topics/{tid}/rewrite", headers=self.h, json={"mode": "simpler"}).status_code, 400)   # no notes yet
        server.db.run(server.update(server.db.topics).where(server.db.topics.c.id == tid).values(notes="## Cells" + chr(10) + "The basic unit."))
        with mock.patch.object(ai, "rewrite_notes", return_value="Simple cells.") as made:
            r = self.c.post(f"/api/topics/{tid}/rewrite", headers=self.h, json={"mode": "translate", "lang": "Arabic"}).get_json()
        self.assertEqual((r["text"], r["lang"], made.call_args.args[2:]), ("Simple cells.", "Arabic", ("translate", "Arabic")))


class Reading(unittest.TestCase):
    def test_text_and_unknown_files(self):
        kind, text = ingest.read_file("notes.txt", ("Photosynthesis uses light. " * 5).encode())
        self.assertEqual(kind, "text")
        with self.assertRaises(ai.AIError):
            ingest.read_file("virus.exe", b"MZ" * 40)
        with self.assertRaises(ai.AIError):
            ingest.read_file("empty.txt", b"hi")

    def test_docx_tables_and_pptx_notes(self):
        import docx
        import pptx
        d = docx.Document()
        d.add_paragraph("Organelles do the work of the cell, each with a job.")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Ribosome", "Makes protein"
        buf = io.BytesIO(); d.save(buf)
        self.assertIn("Ribosome | Makes protein", ingest.read_file("a.docx", buf.getvalue())[1])
        p = pptx.Presentation()
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text, s.placeholders[1].text = "Newton", "F = ma tells us force is mass times acceleration."
        s.notes_slide.notes_text_frame.text = "Mention net force."
        buf = io.BytesIO(); p.save(buf)
        self.assertIn("net force", ingest.read_file("a.pptx", buf.getvalue())[1].lower())

    def test_retrieval_prefers_relevant_chunks(self):
        text = ("Bananas are yellow fruit grown in warm places. " * 80) + ("Mitochondria produce ATP by cellular respiration. " * 5) + ("Cats sleep a lot every day. " * 80)
        got = ai.relevant(text, "How do mitochondria make ATP?", limit=1500, size=400)
        self.assertIn("Mitochondria", got)
        self.assertLess(len(got), len(text))


if __name__ == "__main__":
    unittest.main()
