"""Turning an uploaded file into plain text: PDF, Word, PowerPoint, text, and photos of notes."""

import os
import re

import ai

MAX_TEXT = 400_000
IMAGES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}


def clean(text):
    text = text.replace("\x00", " ").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


import time
import zipfile

MAX_FILE = 1024 * 1024 * 1024                       # 1 GB per file; text is read in a stream, never the whole file at once
MAX_IMAGE = 20 * 1024 * 1024                        # the vision model's own limit
MAX_SCAN = 12 * 1024 * 1024                         # scanned PDFs are read by the AI, so only smallish ones
MAX_XML = 150 * 1024 * 1024                         # text inside a Word/PowerPoint file (guards against zip bombs)
PDF_SECONDS = 70


def _pdf(path, name):
    from pypdf import PdfReader
    started, parts, total = time.time(), [], 0
    with open(path, "rb") as fh:
        reader = PdfReader(fh)
        for page in reader.pages:                   # pages are read one at a time, and we stop once there is enough text
            t = page.extract_text() or ""
            parts.append(t)
            total += len(t)
            if total >= MAX_TEXT * 1.2 or time.time() - started > PDF_SECONDS:
                break
    text = "\n\n".join(parts)
    if len(clean(text)) < 200:                      # scanned pages: have the AI read them
        if os.path.getsize(path) > MAX_SCAN:
            raise ai.AIError(f"{name} looks like scanned pages, and scans can be up to 12 MB. Try fewer pages or a text version.")
        with open(path, "rb") as fh:
            text = ai.pdf_text(fh.read(), name)
    return text


def _zip_guard(path, name):
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist() if i.filename.endswith((".xml", ".rels"))) > MAX_XML:
            raise ai.AIError(f"{name} has far too much text inside. Try splitting it up.")


def _docx(path):
    import docx
    d = docx.Document(path)
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _pptx(path):
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                parts.extend(" | ".join(c.text for c in r.cells) for r in shape.table.rows)
        if slide.has_notes_slide:
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
        out.append(f"[Slide {i}]\n" + "\n".join(p for p in parts if p.strip()))
        if sum(len(p) for p in out) > MAX_TEXT * 1.2:
            break
    return "\n\n".join(out)


def read_path(name, path):
    """Returns (kind, text) from a file on disk. Raises ai.AIError with a friendly message when it can't."""
    ext = os.path.splitext(name.lower())[1]
    size = os.path.getsize(path)
    if size > MAX_FILE:
        raise ai.AIError(f"{name} is bigger than 1 GB. Try a smaller file.")
    try:
        if ext == ".pdf":
            kind, text = "pdf", _pdf(path, name)
        elif ext == ".docx":
            _zip_guard(path, name)
            kind, text = "docx", _docx(path)
        elif ext == ".pptx":
            _zip_guard(path, name)
            kind, text = "pptx", _pptx(path)
        elif ext in IMAGES:
            if size > MAX_IMAGE:
                raise ai.AIError(f"{name} is a photo bigger than 20 MB. Try a smaller one.")
            with open(path, "rb") as fh:
                kind, text = "image", ai.image_text(fh.read(), IMAGES[ext])
        elif ext in (".txt", ".md", ".csv", ".srt", ".vtt", ""):
            with open(path, "rb") as fh:
                kind, text = "text", fh.read(MAX_TEXT * 4).decode("utf-8", "replace")
        else:
            raise ai.AIError(f"{name}: I can read PDF, Word, PowerPoint, text files and photos.")
    except ai.AIError:
        raise
    except Exception:
        raise ai.AIError(f"I couldn't open {name}. Is it a normal, unlocked file?")
    text = clean(text)
    if len(text) < 40:
        raise ai.AIError(f"I couldn't find any text in {name}.")
    return kind, text[:MAX_TEXT]


def read_file(name, data):
    """The same from bytes already in memory (small files, tests)."""
    import tempfile
    fd, path = tempfile.mkstemp(prefix="cramly_up_")
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        return read_path(name, path)
    finally:
        os.remove(path)
